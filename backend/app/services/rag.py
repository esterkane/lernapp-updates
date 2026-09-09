"""Learner-documents RAG layer (ADR-0007, ADR-0008).

Pipeline: ``ingest`` → extraction (pypdf / python-docx / plain text) → German-aware normalisation
(hyphenation repair at line breaks, heading detection) → sentence-safe chunking (~400 tokens,
60-token overlap) → metered embeddings → ``documents`` + ``chunks`` rows.

Retrieval: ``search`` runs a lexical query (``websearch_to_tsquery('german')`` + ``ts_rank_cd``)
and a vector query (pgvector cosine distance, restricted to chunks embedded with the *current*
model so dimensions never mix) and fuses both with pure-Python Reciprocal Rank Fusion.
Every query is scoped by ``owner_id`` (ADR-0008).

CSV uploads are vocabulary lists and are handed to ``app.services.vocab`` without embedding.
"""

from __future__ import annotations

import io
import logging
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from contextvars import copy_context
from dataclasses import dataclass
from pathlib import PurePath
from time import perf_counter
from typing import Any

from pydantic import BaseModel
from sqlalchemy import Select, delete, func, select

from app.core import ledger
from app.core.config import get_settings
from app.core.db import db_session
from app.core.hooks import HookDenied, SessionState, TierCall, engine
from app.db.base import Chunk, Document, VocabItem
from app.services.embeddings import get_embeddings
from app.services.query_cache import query_vector

log = logging.getLogger(__name__)

# ~400 tokens at ≈4 chars/token; 60-token overlap.
CHUNK_CHARS = 1600
OVERLAP_CHARS = 240
TOP_N = 20  # candidates per retriever before fusion
RRF_K = 60
USE_IN_KEYS: tuple[str, ...] = ("uebungen", "rollenspiel", "tutor")
STATUS_INDEXED = "indexiert"
STATUS_ERROR = "fehler"

MIME_BY_EXT: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
}


class ChunkHit(BaseModel):
    chunk_id: str
    document_id: str
    document_title: str
    ord: int
    heading: str | None = None
    text: str
    score: float = 0.0


@dataclass
class ChunkDraft:
    ord: int
    heading: str | None
    text: str


# ---------------------------------------------------------------- parsing helpers


def parse_tags(value: str | list[str] | None) -> list[str]:
    if value is None:
        return []
    raw = value.split(",") if isinstance(value, str) else value
    out: list[str] = []
    for t in raw:
        tag = t.strip().lower()
        if tag and tag not in out:
            out.append(tag)
    return out


def parse_use_in(value: str | list[str] | dict[str, Any] | None) -> dict[str, bool]:
    """``"uebungen,tutor"`` / ``["tutor"]`` / ``{"tutor": true}`` → ``{uebungen, rollenspiel, tutor}`` bools.

    ``None`` means "use everywhere" (the sensible default for a non-technical user).
    """
    if value is None:
        return dict.fromkeys(USE_IN_KEYS, True)
    if isinstance(value, dict):
        unknown = set(value) - set(USE_IN_KEYS)
        if unknown:
            raise ValueError(f"Unbekannte Verwendung: {', '.join(sorted(unknown))}")
        return {k: bool(value.get(k, False)) for k in USE_IN_KEYS}
    keys = parse_tags(value)
    unknown = set(keys) - set(USE_IN_KEYS)
    if unknown:
        raise ValueError(f"Unbekannte Verwendung: {', '.join(sorted(unknown))}")
    return {k: k in keys for k in USE_IN_KEYS}


def mime_for(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in MIME_BY_EXT:
        raise ValueError(f"Nicht unterstütztes Dateiformat: {ext or filename} (erlaubt: PDF, DOCX, TXT, MD, CSV)")
    return MIME_BY_EXT[ext]


def title_for(filename: str) -> str:
    stem = PurePath(filename).stem.replace("_", " ").replace("-", " ").strip()
    return (stem or filename)[:255]


# ---------------------------------------------------------------- extraction


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None and p.style.name else "").lower()
        if style.startswith("heading") or style == "title":
            parts.append(f"# {text}")
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(c for c in cells if c))
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext in (".txt", ".md"):
        return _decode(data)
    raise ValueError(f"Nicht unterstütztes Dateiformat: {ext or filename}")


# ---------------------------------------------------------------- normalisation & structure

_LOWER = "a-zäöüß"
_UPPER = "A-ZÄÖÜ"
# `Ver-\nhandlung` → `Verhandlung`; a hyphen followed by an upper-case start (`E-\nMail`) is a real hyphen.
_HYPHEN_JOIN = re.compile(rf"(\w)-[ \t]*\n[ \t]*(?=[{_LOWER}])")
_HYPHEN_KEEP = re.compile(rf"(\w)-[ \t]*\n[ \t]*(?=[{_UPPER}0-9])")
_BULLET = re.compile(r"^\s*([-*•–]|\d+[.)])\s+")
# fmt: off
_ABBREVIATIONS = frozenset((
    "z", "b", "bzw", "dr", "prof", "nr", "s", "u", "a", "ca", "vgl", "usw", "etc", "d", "h", "evtl",
    "ggf", "inkl", "max", "min", "mio", "mrd", "sog", "str", "tel", "zb", "abs", "art", "bsp", "geb",
    "jh", "o", "ä",
))
# fmt: on
_SENTENCE_END = re.compile(rf"([.!?…]+[\"“”'»«)\]]*)\s+(?=[{_UPPER}0-9„\"“»(\[])")


def repair_hyphenation(text: str) -> str:
    text = text.replace("­", "")  # soft hyphen
    text = _HYPHEN_JOIN.sub(r"\1", text)
    return _HYPHEN_KEEP.sub(r"\1-", text)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = repair_hyphenation(text)
    text = re.sub(r"[  ]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def is_heading(line: str, *, prev_blank: bool) -> bool:
    s = line.strip()
    if not s:
        return False
    if s.startswith("#"):
        return True
    letters = [c for c in s if c.isalpha()]
    if len(letters) >= 2 and s.isupper() and len(s) <= 100:
        return True
    if not prev_blank or _BULLET.match(s):
        return False
    return len(s) <= 60 and len(s.split()) <= 8 and s[-1] not in ".!?:;,"


def _heading_text(line: str) -> str:
    return line.strip().lstrip("#").strip()[:255]


def split_sentences(paragraph: str) -> list[str]:
    """Split on sentence-final punctuation followed by an upper-case start; guards common abbreviations."""
    out: list[str] = []
    start = 0
    for m in _SENTENCE_END.finditer(paragraph):
        candidate = paragraph[start : m.end(1)]
        last_word = re.split(r"\s+", candidate.rstrip(".!?…\"“”'»«)]"))[-1].lower().strip('(„"“')
        if candidate.rstrip().endswith(".") and (last_word in _ABBREVIATIONS or len(last_word) == 1):
            continue
        out.append(candidate.strip())
        start = m.end()
    tail = paragraph[start:].strip()
    if tail:
        out.append(tail)
    return [s for s in out if s]


def structure(text: str) -> list[tuple[str | None, list[str]]]:
    """Return ``[(heading, [paragraph, ...]), ...]`` with hard-wrapped lines joined into paragraphs."""
    sections: list[tuple[str | None, list[str]]] = []
    heading: str | None = None
    paragraphs: list[str] = []
    buf: list[str] = []
    prev_blank = True

    def flush() -> None:
        if buf:
            paragraphs.append(" ".join(buf).strip())
            buf.clear()

    for line in normalize_text(text).split("\n"):
        if not line.strip():
            flush()
            prev_blank = True
            continue
        if is_heading(line, prev_blank=prev_blank):
            flush()
            if paragraphs or heading is not None:
                sections.append((heading, paragraphs))
                paragraphs = []
            heading = _heading_text(line)
        elif _BULLET.match(line):
            flush()
            buf.append(line.strip())
        else:
            buf.append(line.strip())
        prev_blank = False
    flush()
    if paragraphs or heading is not None:
        sections.append((heading, paragraphs))
    return [(h, [p for p in ps if p]) for h, ps in sections if ps]


# ---------------------------------------------------------------- chunking


def chunk_text(text: str, *, max_chars: int = CHUNK_CHARS, overlap_chars: int = OVERLAP_CHARS) -> list[ChunkDraft]:
    """Sentence-safe chunks: split on headings/paragraphs first, then sentences, never mid-sentence.

    Consecutive chunks of one section overlap by whole trailing sentences (≤ ``overlap_chars``).
    A heading is repeated as the first line of every chunk of its section so lexical and vector
    retrieval can see it.
    """
    drafts: list[ChunkDraft] = []

    def emit(heading: str | None, units: list[tuple[str, int]]) -> None:
        if not units:
            return
        body_parts: list[str] = []
        cur_par = units[0][1]
        cur: list[str] = []
        for sentence, par in units:
            if par != cur_par:
                body_parts.append(" ".join(cur))
                cur, cur_par = [], par
            cur.append(sentence)
        body_parts.append(" ".join(cur))
        body = "\n\n".join(body_parts)
        drafts.append(ChunkDraft(len(drafts), heading, f"{heading}\n{body}" if heading else body))

    for heading, paragraphs in structure(text):
        units: list[tuple[str, int]] = []
        size = len(heading or "")
        for par_idx, paragraph in enumerate(paragraphs):
            for sentence in split_sentences(paragraph):
                if units and size + len(sentence) + 1 > max_chars:
                    emit(heading, units)
                    keep: list[tuple[str, int]] = []
                    kept = 0
                    for s, p in reversed(units):
                        if kept + len(s) > overlap_chars:
                            break
                        keep.insert(0, (s, p))
                        kept += len(s) + 1
                    units = keep
                    size = len(heading or "") + kept
                units.append((sentence, par_idx))
                size += len(sentence) + 1
        emit(heading, units)
    return drafts


# ---------------------------------------------------------------- ingestion


def _doc_dict(d: Document) -> dict[str, Any]:
    return {
        "id": d.id,
        "title": d.title,
        "filename": d.filename,
        "mime": d.mime,
        "tags": list(d.tags or []),
        "use_in": {k: bool((d.use_in or {}).get(k, False)) for k in USE_IN_KEYS},
        "n_chunks": d.n_chunks,
        "n_chars": d.n_chars,
        "status": d.status,
        "error": d.error,
        "created_at": d.created_at.isoformat() if d.created_at else None,
    }


def _set_status(doc_id: str, status: str, **fields: Any) -> None:
    with db_session() as db:
        doc = db.get(Document, doc_id)
        if doc is None:
            return
        doc.status = status
        for k, v in fields.items():
            setattr(doc, k, v)


def ingest(
    learner_id: str,
    filename: str,
    data: bytes,
    tags: str | list[str] | None = None,
    use_in: str | list[str] | dict[str, Any] | None = None,
    *,
    title: str | None = None,
) -> dict[str, Any]:
    """Store + index one uploaded file.

    Returns ``{document_id, title, status, error, n_chunks, n_vocab_items, cost_eur}``.

    Status progression on ``documents.status``: hochgeladen → extrahiert → gechunkt → indexiert,
    or ``fehler`` with ``error`` text. Unsupported file types raise ``ValueError`` before a row exists.
    """
    mime = mime_for(filename)
    tag_list = parse_tags(tags)
    use_in_map = parse_use_in(use_in)
    with db_session() as db:
        doc = Document(
            owner_id=learner_id,
            title=(title or title_for(filename))[:255],
            filename=filename[:255],
            mime=mime,
            tags=tag_list,
            use_in=use_in_map,
            status="hochgeladen",
        )
        db.add(doc)
        db.flush()
        doc_id = doc.id
        doc_title = doc.title
    cost_key = f"doc:{doc_id}"
    n_chunks = 0
    n_vocab = 0
    try:
        if mime == "text/csv":
            from app.services import vocab

            n_vocab = vocab.import_csv(learner_id, data, source_document_id=doc_id)
            if n_vocab == 0 and not vocab.parse_csv(data):
                raise ValueError("Keine Vokabeln gefunden (erwartet: wort,bedeutung[,beispiel])")
            _set_status(doc_id, STATUS_INDEXED, n_chars=len(data), n_chunks=0, error=None)
        else:
            text = extract_text(filename, data)
            text = normalize_text(text)
            if not text.strip():
                raise ValueError("Kein Text in der Datei gefunden")
            _set_status(doc_id, "extrahiert", n_chars=len(text))
            drafts = chunk_text(text)
            if not drafts:
                raise ValueError("Kein Text in der Datei gefunden")
            _set_status(doc_id, "gechunkt", n_chunks=len(drafts))
            backend = get_embeddings()
            vectors = backend.embed([d.text for d in drafts], session_id=cost_key, learner_id=learner_id)
            with db_session() as db:
                for d, vec in zip(drafts, vectors, strict=True):
                    db.add(
                        Chunk(
                            document_id=doc_id,
                            owner_id=learner_id,
                            ord=d.ord,
                            heading=d.heading,
                            text=d.text,
                            embedding=vec,
                            embedding_model=backend.model,
                        )
                    )
            n_chunks = len(drafts)
            _set_status(doc_id, STATUS_INDEXED, n_chunks=n_chunks, error=None)
        status = STATUS_INDEXED
        error: str | None = None
    except Exception as exc:  # noqa: BLE001 — recorded on the row, surfaced to the UI
        log.warning("ingest failed for %s: %s", filename, exc)
        status, error = STATUS_ERROR, str(exc)[:2000]
        _set_status(doc_id, STATUS_ERROR, error=error)
    return {
        "document_id": doc_id,
        "title": doc_title,
        "status": status,
        "error": error,
        "n_chunks": n_chunks,
        "n_vocab_items": n_vocab,
        "cost_eur": ledger.session_cost_eur(cost_key),
    }


# ---------------------------------------------------------------- document management


def list_documents(owner_id: str) -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute(
            select(Document).where(Document.owner_id == owner_id).order_by(Document.created_at.desc(), Document.id)
        ).scalars()
        return [_doc_dict(d) for d in rows]


def get_document(document_id: str, owner_id: str | None = None) -> dict[str, Any] | None:
    with db_session() as db:
        doc = db.get(Document, document_id)
        if doc is None or (owner_id is not None and doc.owner_id != owner_id):
            return None
        return _doc_dict(doc)


def update_document(
    document_id: str,
    *,
    owner_id: str | None = None,
    tags: str | list[str] | None = None,
    use_in: str | list[str] | dict[str, Any] | None = None,
    title: str | None = None,
) -> dict[str, Any] | None:
    with db_session() as db:
        doc = db.get(Document, document_id)
        if doc is None or (owner_id is not None and doc.owner_id != owner_id):
            return None
        if tags is not None:
            doc.tags = parse_tags(tags)
        if use_in is not None:
            doc.use_in = parse_use_in(use_in)
        if title is not None and title.strip():
            doc.title = title.strip()[:255]
        db.flush()
        return _doc_dict(doc)


def delete_document(document_id: str, owner_id: str | None = None) -> bool:
    """Delete a document, its chunks and vocab items imported from it."""
    with db_session() as db:
        doc = db.get(Document, document_id)
        if doc is None or (owner_id is not None and doc.owner_id != owner_id):
            return False
        db.execute(delete(Chunk).where(Chunk.document_id == document_id))
        db.execute(delete(VocabItem).where(VocabItem.source_document_id == document_id))
        db.delete(doc)
    return True


# ---------------------------------------------------------------- retrieval


def rrf(rank_lists: list[list[str]], k: int = RRF_K) -> list[tuple[str, float]]:
    """Reciprocal Rank Fusion: score(id) = Σ 1/(k + rank). Deterministic: ties broken by id."""
    scores: dict[str, float] = {}
    for ranked in rank_lists:
        for rank, item in enumerate(ranked, start=1):
            scores[item] = scores.get(item, 0.0) + 1.0 / (k + rank)
    return sorted(scores.items(), key=lambda kv: (-kv[1], kv[0]))


def _scoped(stmt: Select[Any], owner_id: str, tags: list[str] | None, use_in: str | None) -> Select[Any]:
    stmt = stmt.join(Document, Document.id == Chunk.document_id).where(
        Chunk.owner_id == owner_id, Document.owner_id == owner_id, Document.status == STATUS_INDEXED
    )
    if tags:
        stmt = stmt.where(Document.tags.overlap(tags))
    if use_in:
        stmt = stmt.where(Document.use_in[use_in].as_boolean().is_(True))
    return stmt


def search(
    query: str,
    owner_id: str,
    tags: str | list[str] | None = None,
    k: int = 8,
    use_in: str | None = None,
    session_id: str | None = None,
    state: SessionState | None = None,
    timings: dict[str, float] | None = None,
) -> list[ChunkHit]:
    """Hybrid search over the owner's indexed chunks (lexical ∪ vector, RRF-fused).

    Runs through the hook engine as a ``rag`` call: ``rag_owner_scope`` denies retrieval without an
    owner (raises ``HookDenied``). The query embedding is attributed to ``session_id`` or, when the
    caller has none, to ``rag-search:<owner>`` so the ledger row is never orphaned.
    """
    sid = session_id or f"rag-search:{owner_id}"
    call = TierCall(
        kind="rag",
        tier="rag",
        model="hybrid-rrf",
        owner_id=owner_id,
        session_id=sid,
        learner_id=owner_id,
        payload_summary={"k": k, "use_in": use_in, "query_chars": len(query or "")},
    )
    res = engine.execute(
        call,
        state or SessionState(),
        lambda _c: _search(query, owner_id, tags=tags, k=k, use_in=use_in, session_id=sid, timings=timings),
    )
    if not res.success:
        raise HookDenied(res)
    return res.result  # type: ignore[no-any-return]


def _search(
    query: str,
    owner_id: str,
    *,
    tags: str | list[str] | None,
    k: int,
    use_in: str | None,
    session_id: str,
    timings: dict[str, float] | None = None,
) -> list[ChunkHit]:
    if not query or not query.strip() or k <= 0:
        return []
    tag_list = parse_tags(tags) or None
    if use_in is not None and use_in not in USE_IN_KEYS:
        raise ValueError(f"Unbekannte Verwendung: {use_in}")
    if get_settings().rag_backend == "es":
        from app.services.rag_backends import get_backend

        return get_backend().search(query, owner_id, tags=tag_list, k=k, use_in=use_in)

    started = perf_counter()
    with db_session() as db:
        eligible = db.execute(_scoped(select(Chunk.id), owner_id, tag_list, use_in).limit(1)).first()
    database_ms = (perf_counter() - started) * 1000
    if eligible is None:
        if timings is not None:
            timings.update(query_embedding=0.0, retrieval_database=database_ms)
        return []

    backend = get_embeddings()

    def embed_query() -> tuple[list[float], float]:
        clock = perf_counter()
        vector = query_vector(backend, query, owner_id=owner_id, session_id=session_id)
        return vector, (perf_counter() - clock) * 1000

    # Each branch owns its DB sessions. Copy request-local workspace/audit context
    # into the embedding worker so credentials, hooks and billing stay scoped.
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix="rag-embedding") as pool:
        pending = pool.submit(copy_context().run, embed_query)
        started = perf_counter()
        tsq = func.websearch_to_tsquery("german", query)
        with db_session() as db:
            lex_stmt = (
                _scoped(select(Chunk.id), owner_id, tag_list, use_in)
                .where(Chunk.tsv.op("@@")(tsq))
                .order_by(func.ts_rank_cd(Chunk.tsv, tsq).desc(), Chunk.id)
                .limit(TOP_N)
            )
            lexical = [str(r) for r in db.execute(lex_stmt).scalars()]
        database_ms += (perf_counter() - started) * 1000
        query_vec, embedding_ms = pending.result()
    if timings is not None:
        timings["query_embedding"] = embedding_ms
    started = perf_counter()
    with db_session() as db:
        distance = Chunk.embedding.cosine_distance(query_vec)
        vec_stmt = (
            _scoped(select(Chunk.id), owner_id, tag_list, use_in)
            .where(Chunk.embedding.is_not(None), Chunk.embedding_model == backend.model)
            .order_by(distance, Chunk.id)
            .limit(TOP_N)
        )
        vector = [str(r) for r in db.execute(vec_stmt).scalars()]
        fused = rrf([lexical, vector])[:k]
        if not fused:
            if timings is not None:
                timings["retrieval_database"] = database_ms + (perf_counter() - started) * 1000
            return []
        ids = [cid for cid, _ in fused]
        rows = db.execute(
            _scoped(select(Chunk, Document.title), owner_id, tag_list, use_in).where(Chunk.id.in_(ids))
        ).all()
    by_id = {c.id: (c, title) for c, title in rows}
    hits: list[ChunkHit] = []
    for cid, score in fused:
        if cid not in by_id:
            continue
        c, title = by_id[cid]
        hits.append(
            ChunkHit(
                chunk_id=c.id,
                document_id=c.document_id,
                document_title=title,
                ord=c.ord,
                heading=c.heading,
                text=c.text,
                score=round(score, 6),
            )
        )
    if timings is not None:
        timings["retrieval_database"] = database_ms + (perf_counter() - started) * 1000
    return hits


def context_for(
    learner_id: str,
    query: str,
    *,
    use_in: str = "tutor",
    tags: list[str] | None = None,
    k: int = 5,
    session_id: str | None = None,
    timings: dict[str, float] | None = None,
) -> list[ChunkHit]:
    """Retrieval helper for tutor / roleplay / task generation. Never raises — retrieval is best-effort."""
    try:
        return search(query, learner_id, tags=tags, k=k, use_in=use_in, session_id=session_id, timings=timings)
    except Exception:  # noqa: BLE001
        log.exception("context_for failed (learner=%s)", learner_id)
        return []


def format_context(hits: list[ChunkHit]) -> str:
    if not hits:
        return "—"
    return "\n\n".join(f"[Dok: {h.document_title} §{h.ord}] {h.text}" for h in hits)
